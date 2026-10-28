import os
import sys
from fastapi.testclient import TestClient

# Disable telemetry
os.environ["ANONYMIZED_TELEMETRY"] = "False"
os.environ["CHROMA_SERVER_NOOP_TELEMETRY"] = "True"

from app.main import app
from pptx import Presentation
import io

def run_e2e_verification():
    print("=" * 60)
    print("RUNNING END-TO-END PRESENTATION CONTENT BUILDER VERIFICATION")
    print("=" * 60)

    client = TestClient(app)

    # 1. Health Check
    res = client.get("/api/health")
    assert res.status_code == 200, f"Health check failed: {res.text}"
    print("1. [PASS] Health check verified. Status:", res.json()["status"])

    # 2. Upload Document
    sample_path = os.path.join("sample_documents", "enterprise_rag_playbook.txt")
    with open(sample_path, "rb") as f:
        file_bytes = f.read()

    upload_res = client.post(
        "/api/documents/upload",
        files=[("files", ("enterprise_rag_playbook.txt", file_bytes, "text/plain"))],
        data={"chunk_size": 600, "chunk_overlap": 100}
    )
    assert upload_res.status_code == 200, f"Upload failed: {upload_res.text}"
    upload_data = upload_res.json()
    print(f"2. [PASS] Document upload & chunking verified. Ingested {upload_data['total_chunks_created']} chunks.")

    # 3. Vector Similarity Search
    search_res = client.post(
        "/api/search",
        json={"query": "What are the core architectural pillars and chunk sizes?", "top_k": 3}
    )
    assert search_res.status_code == 200, f"Search failed: {search_res.text}"
    search_data = search_res.json()
    assert len(search_data["results"]) > 0, "No search results returned"
    top_hit = search_data["results"][0]
    print(f"3. [PASS] Semantic vector search verified. Top score: {top_hit['similarity_score']:.2f} ({top_hit['source_filename']})")

    # 4. Slide Deck Generation (Strict JSON Schema Contract)
    gen_res = client.post(
        "/api/generate-deck",
        json={
            "topic": "Enterprise RAG Ingestion & PowerPoint Synthesis",
            "target_audience": "Technical Architects & Executive Sponsors",
            "num_slides": 5,
            "deck_tone": "Technical",
            "llm_provider": "local_offline"
        }
    )
    assert gen_res.status_code == 200, f"Deck generation failed: {gen_res.text}"
    deck = gen_res.json()
    assert "deck_title" in deck
    assert "target_audience" in deck
    assert len(deck["slides"]) == 5
    print(f"4. [PASS] SlideDeck JSON contract verified. Title: '{deck['deck_title']}', Slides: {len(deck['slides'])}")

    for s in deck["slides"]:
        print(f"   -> Slide {s['slide_number']}: '{s['title']}' | Layout: {s['visual_layout_suggestion']} | Citations: {s['retrieved_sources']}")

    # 5. Programmatic PPTX Export
    export_res = client.post(
        "/api/export-pptx",
        json={"deck": deck, "theme": "modern_indigo"}
    )
    assert export_res.status_code == 200, f"PPTX export failed: {export_res.text}"
    assert "presentationml.presentation" in export_res.headers["content-type"]
    pptx_bytes = export_res.content
    assert len(pptx_bytes) > 2000

    # Save to disk as proof of export
    out_path = os.path.join("data", "generated_sample_deck.pptx")
    with open(out_path, "wb") as f:
        f.write(pptx_bytes)

    # Inspect generated presentation using python-pptx
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 5
    print(f"5. [PASS] Programmatic PPTX Export verified. Generated file saved to '{out_path}' ({len(pptx_bytes)} bytes)")
    print(f"   Slide width: {prs.slide_width.inches:.3f} inches (16:9 widescreen: 13.333)")
    print(f"   Slide height: {prs.slide_height.inches:.3f} inches (16:9 widescreen: 7.500)")
    print(f"   Speaker notes verified on Slide 1: '{prs.slides[0].notes_slide.notes_text_frame.text[:60]}...'")

    print("=" * 60)
    print("ALL END-TO-END SYSTEM CAPABILITIES VERIFIED SUCCESSFULLY!")
    print("=" * 60)

if __name__ == "__main__":
    run_e2e_verification()


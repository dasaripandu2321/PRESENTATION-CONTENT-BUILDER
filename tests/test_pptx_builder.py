import io
import pytest
from pptx import Presentation
from app.services.pptx_builder import PPTXBuilder
from app.models.schemas import SlideDeck, Slide

def test_pptx_generation():
    builder = PPTXBuilder()
    
    deck = SlideDeck(
        deck_title="AI Presentation Builder Test",
        target_audience="Engineers & Leadership",
        slides=[
            Slide(
                slide_number=1,
                title="Title Slide Hero",
                core_takeaway="Executive takeaway statement.",
                bullet_points=["Key bullet point 1", "Key bullet point 2"],
                visual_layout_suggestion="title_hero",
                speaker_notes="Presenter talking point script.",
                retrieved_sources=["spec.pdf (p. 1)"]
            ),
            Slide(
                slide_number=2,
                title="Split 2 Column Analysis",
                core_takeaway="Secondary takeaway statement.",
                bullet_points=["Column bullet A", "Column bullet B"],
                visual_layout_suggestion="split_2_column",
                speaker_notes="Second slide notes.",
                retrieved_sources=["spec.pdf (p. 2)"]
            ),
            Slide(
                slide_number=3,
                title="3 Pillar Strategy",
                core_takeaway="Pillar takeaway statement.",
                bullet_points=["Pillar 1", "Pillar 2", "Pillar 3"],
                visual_layout_suggestion="3_pillar_cards",
                speaker_notes="Pillar slide notes.",
                retrieved_sources=["spec.pdf (p. 3)"]
            ),
            Slide(
                slide_number=4,
                title="KPI Metrics Highlight",
                core_takeaway="KPI takeaway statement.",
                bullet_points=["KPI Metric 1", "KPI Metric 2"],
                visual_layout_suggestion="kpi_highlight",
                speaker_notes="KPI slide notes.",
                retrieved_sources=["spec.pdf (p. 4)"]
            )
        ]
    )

    pptx_bytes = builder.build_presentation(deck, theme_name="modern_indigo")
    assert isinstance(pptx_bytes, io.BytesIO)
    assert pptx_bytes.getbuffer().nbytes > 1000

    # Verify generated presentation with pptx.Presentation
    prs = Presentation(pptx_bytes)
    assert len(prs.slides) == 4
    # Verify 16:9 aspect ratio (13.333 inches x 7.5 inches)
    assert round(prs.slide_width.inches, 2) == 13.33
    assert round(prs.slide_height.inches, 2) == 7.5

    # Verify speaker notes attached
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Presenter talking point script."

